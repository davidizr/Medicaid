import io
import re
import warnings
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import statsmodels.api as sm
from statsmodels.tools.sm_exceptions import ConvergenceWarning
from matplotlib.ticker import FuncFormatter
from pptx import Presentation
from pptx.util import Inches, Pt
from scipy.stats import chi2


EMU_PER_PT = 12700
UTILIZATION_RATE_COLUMNS = [
    "code",
    "period_type",
    "period_label",
    "claims",
    "exposure_months",
    "rate_per_100000_member_months",
    "poisson_ci_lower_per_100000",
    "poisson_ci_upper_per_100000",
]


AGG_MAP = {
    "Agg_autonomic_testing": ["95921", "95922", "95923", "95924"],
    "Agg_skin_biopsy": ["11100", "11101", "11102", "11103", "11104", "11105", "11106", "11107"],
    "Agg_cardiac_monitoring": [
        "0295T",
        "0296T",
        "0297T",
        "0298T",
        "93241",
        "93242",
        "93243",
        "93244",
        "93245",
        "93246",
        "93247",
        "93248",
    ],
    "Agg_fractures": [
        "23500",
        "23600",
        "25600",
        "27500",
        "27750",
        "27786",
        "28470",
        "21310",
        "28510",
        "25605",
        "24670",
        "27236",
        "24546",
        "25606",
        "27235",
    ],
    "Agg_nerve_conduction": ["95907", "95908", "95909", "95910", "95911", "95912", "95913"],
    "Agg_neurophysiologic/autonomic testing": ["95943", "95999"],
    "Agg_immunohistochemistry_stains": ["88341", "88342"],
    "Agg_psychological_testing": [
        "96130",
        "96131",
        "96132",
        "96133",
        "96134",
        "96135",
        "96136",
        "96137",
        "96138",
        "96139",
        "96146",
    ],
    "Agg_neurobehavioral": ["96116", "96121"],
    "Agg_needle_EMG": ["95886", "95887"],
    "Agg_clotting": [
        "85240",
        "85241",
        "85242",
        "85243",
        "85244",
        "85245",
        "85246",
        "85247",
        "85248",
        "85249",
        "85250",
        "85251",
        "85252",
        "85253",
        "85254",
        "85255",
        "85256",
        "85257",
        "85258",
        "85259",
        "85260",
        "85261",
        "85262",
        "85263",
        "85264",
        "85265",
        "85266",
        "85267",
        "85268",
        "85269",
        "85270",
    ],
    "Agg_pregnancy_test": ["81025", "84073"],
    "Agg_fluorescent_noninfectious_agent": ["86383", "86051", "86052", "86053", "86255", "86256"],
}


def load_claims_pivot(path):
    claims = pd.read_csv(path)
    date_col = claims.columns[0]
    claims[date_col] = pd.to_datetime(claims[date_col])
    claims = claims.set_index(date_col).sort_index()
    return ensure_aggregates(claims)


def load_enrollment(path):
    enrollment = pd.read_excel(path)
    enrollment = enrollment.rename(columns={col: str(col).strip() for col in enrollment.columns})
    if {"Date", "Total Enrollees"}.issubset(enrollment.columns):
        enrollment = enrollment[["Date", "Total Enrollees"]].rename(
            columns={"Date": "month", "Total Enrollees": "enrollment"}
        )
    elif {"Month", "Enrollment"}.issubset(enrollment.columns):
        enrollment = enrollment[["Month", "Enrollment"]].rename(
            columns={"Month": "month", "Enrollment": "enrollment"}
        )
    else:
        raise ValueError("Enrollment file must contain Date/Total Enrollees or Month/Enrollment columns.")

    enrollment["month"] = pd.to_datetime(enrollment["month"]).dt.to_period("M").dt.to_timestamp()
    enrollment["enrollment"] = pd.to_numeric(enrollment["enrollment"], errors="coerce")
    enrollment = enrollment.dropna(subset=["month", "enrollment"])
    enrollment = enrollment[enrollment["enrollment"] > 0]
    enrollment["enrollment"] = enrollment["enrollment"].astype(float)
    enrollment = enrollment.drop_duplicates(subset=["month"], keep="last").sort_values("month")

    interpolate_month = pd.Timestamp("2022-04-01")
    prev_month = pd.Timestamp("2022-03-01")
    next_month = pd.Timestamp("2022-05-01")
    month_values = enrollment.set_index("month")["enrollment"]
    if {interpolate_month, prev_month, next_month}.issubset(month_values.index):
        interpolated = (month_values.loc[prev_month] + month_values.loc[next_month]) / 2
        enrollment.loc[enrollment["month"] == interpolate_month, "enrollment"] = interpolated

    return enrollment


def ensure_aggregates(claims):
    out = claims.copy()
    lookup = {str(col).strip().upper(): col for col in out.columns}
    for agg_name, codes in AGG_MAP.items():
        present = [lookup[str(code).strip().upper()] for code in codes if str(code).strip().upper() in lookup]
        out[agg_name] = out[present].fillna(0).sum(axis=1) if present else 0.0
    return out


def norm_text(value):
    text = str(value).strip().lower()
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def norm_agg_code(value):
    text = re.sub(r"(?i)^agg[_\s-]*", "", str(value).strip())
    return norm_text(text)


def aggregate_title_lookup(claims):
    lookup = {}
    for agg_col in AGG_MAP:
        title = "Aggregate " + norm_agg_code(agg_col).title()
        lookup[norm_text(title)] = agg_col
    for col in claims.columns:
        col_text = str(col).strip()
        if col_text.upper().startswith("AGG"):
            title = "Aggregate " + norm_agg_code(col_text).title()
            lookup[norm_text(title)] = col_text
    return lookup


def slide_title(slide):
    try:
        if slide.shapes.title is not None and slide.shapes.title.text:
            return slide.shapes.title.text.strip()
    except Exception:
        pass
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False) and shape.text and shape.text.strip():
            return shape.text.strip().splitlines()[0]
    return ""


def slide_code(title, claims, agg_lookup):
    if title.lower().startswith("normalized its"):
        return None

    hcpcs = re.match(r"^\s*HCPCS\s+([A-Z0-9]+)\b", title, flags=re.IGNORECASE)
    if hcpcs:
        code = hcpcs.group(1).strip().upper()
        claims_lookup = {str(col).strip().upper(): str(col).strip() for col in claims.columns}
        return claims_lookup.get(code)

    aggregate = re.match(r"^\s*Aggregate\s+(.+?)\s*$", title, flags=re.IGNORECASE)
    if aggregate:
        key = norm_text("Aggregate " + aggregate.group(1).strip())
        return agg_lookup.get(key)

    return None


def extract_template_geometry(slide):
    geom = {
        "title_left": 177800,
        "title_top": 228600,
        "title_width": 11836400,
        "title_height": 461665,
        "plot_left": 896555,
        "plot_top": 1646674,
        "plot_width": 5395090,
        "plot_height": 4275853,
        "summary_left": 7137400,
        "summary_top": 889000,
        "summary_width": 4376198,
        "summary_height": 2954655,
    }

    text_shapes = [shape for shape in slide.shapes if getattr(shape, "has_text_frame", False)]
    pic_shapes = [shape for shape in slide.shapes if shape.shape_type == 13]

    if pic_shapes:
        pic = pic_shapes[0]
        geom.update(
            {
                "plot_left": int(pic.left),
                "plot_top": int(pic.top),
                "plot_width": int(pic.width),
                "plot_height": int(pic.height),
            }
        )

    if text_shapes:
        text_shapes = sorted(text_shapes, key=lambda shape: (int(shape.top), int(shape.left)))
        title_shape = text_shapes[0]
        geom.update(
            {
                "title_left": int(title_shape.left),
                "title_top": int(title_shape.top),
                "title_width": int(title_shape.width),
                "title_height": int(title_shape.height),
            }
        )
        summary_candidates = [
            shape for shape in text_shapes[1:] if int(shape.width) > 0 and int(shape.height) > 0
        ]
        if summary_candidates:
            summary_shape = max(summary_candidates, key=lambda shape: int(shape.width) * int(shape.height))
            geom.update(
                {
                    "summary_left": int(summary_shape.left),
                    "summary_top": int(summary_shape.top),
                    "summary_width": int(summary_shape.width),
                    "summary_height": int(summary_shape.height),
                }
            )

    return geom


def trim_to_first_nonzero(ts):
    ts = ts.copy()
    ts["month"] = pd.to_datetime(ts["month"])
    ts["y"] = pd.to_numeric(ts["y"], errors="coerce")
    nonzero_idx = ts.index[ts["y"].fillna(0) > 0]
    if len(nonzero_idx) == 0:
        return ts.iloc[0:0].copy()
    return ts.loc[nonzero_idx[0] :].reset_index(drop=True)


def build_ts(claims, code_col):
    ts = claims[[code_col]].reset_index()
    ts.columns = ["month", "y"]
    ts["month"] = pd.to_datetime(ts["month"]).dt.to_period("M").dt.to_timestamp()
    ts["y"] = pd.to_numeric(ts["y"], errors="coerce")
    return trim_to_first_nonzero(ts)


def add_enrollment_and_its_terms(ts_raw, enrollment, intervention_month):
    ts = trim_to_first_nonzero(ts_raw)
    if ts.empty:
        raise ValueError("series has no non-zero observations")

    ts["y"] = ts["y"].fillna(0)
    enrollment = enrollment.copy()
    enrollment["month"] = pd.to_datetime(enrollment["month"]).dt.to_period("M").dt.to_timestamp()
    ts = ts.merge(enrollment[["month", "enrollment"]], on="month", how="left", validate="many_to_one")
    if ts["enrollment"].isna().any():
        missing = ts.loc[ts["enrollment"].isna(), "month"].dt.strftime("%Y-%m").tolist()
        raise ValueError(f"missing enrollment values for months: {missing}")
    if (ts["enrollment"] <= 0).any():
        bad = ts.loc[ts["enrollment"] <= 0, "month"].dt.strftime("%Y-%m").tolist()
        raise ValueError(f"non-positive enrollment values for months: {bad}")

    ts["_month_index"] = ts["month"].dt.year * 12 + ts["month"].dt.month
    ts["T"] = ts["_month_index"] - ts["_month_index"].iloc[0]

    intervention_month = pd.Timestamp(intervention_month).to_period("M").to_timestamp()
    intervention_idx = intervention_month.year * 12 + intervention_month.month
    ts["X_t"] = (ts["month"] >= intervention_month).astype(int)
    ts["post_time"] = 0
    post_mask = ts["month"] >= intervention_month
    if post_mask.any():
        ts.loc[post_mask, "post_time"] = ts.loc[post_mask, "_month_index"] - intervention_idx + 1

    x = sm.add_constant(ts[["T", "X_t", "post_time"]])
    return ts, x


def fit_normalized_poisson(ts_raw, enrollment, intervention_month, rate_per=100000):
    ts, x = add_enrollment_and_its_terms(ts_raw, enrollment, intervention_month)
    model = sm.GLM(ts["y"], x, family=sm.families.Poisson(), exposure=ts["enrollment"]).fit()
    ts["y_hat"] = model.predict(x, exposure=ts["enrollment"])
    ts["rate"] = ts["y"] / ts["enrollment"] * rate_per
    ts["rate_hat"] = ts["y_hat"] / ts["enrollment"] * rate_per
    return ts, model


def fit_normalized_negative_binomial(ts, intervention_month, rate_per=100000):
    fitted_ts, x = add_enrollment_and_its_terms(ts[["month", "y"]], ts[["month", "enrollment"]], intervention_month)
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", category=ConvergenceWarning)
        warnings.filterwarnings("ignore", category=RuntimeWarning)
        model = sm.NegativeBinomial(
            fitted_ts["y"],
            x,
            exposure=fitted_ts["enrollment"],
            loglike_method="nb2",
        ).fit(disp=False, maxiter=200)
    fitted_ts["y_hat"] = model.predict(x, exposure=fitted_ts["enrollment"])
    fitted_ts["rate"] = fitted_ts["y"] / fitted_ts["enrollment"] * rate_per
    fitted_ts["rate_hat"] = fitted_ts["y_hat"] / fitted_ts["enrollment"] * rate_per
    return fitted_ts, model


def exact_poisson_count_ci(count, alpha=0.05):
    count = int(count)
    if count < 0:
        raise ValueError("Poisson confidence interval requires non-negative claims.")
    lower = 0.0 if count == 0 else 0.5 * chi2.ppf(alpha / 2, 2 * count)
    upper = 0.5 * chi2.ppf(1 - alpha / 2, 2 * (count + 1))
    return float(lower), float(upper)


def _period_rate_row(code, period_type, period_label, claims, exposure_months, rate_per=100000):
    claims = float(claims)
    if not np.isclose(claims, round(claims)):
        raise ValueError(f"Poisson confidence interval requires integer claims; got {claims}.")
    claims = int(round(claims))
    exposure_months = float(exposure_months)
    if exposure_months <= 0:
        raise ValueError(f"Exposure months must be positive for {code} {period_label}.")

    ci_lower_count, ci_upper_count = exact_poisson_count_ci(claims)
    return {
        "code": str(code),
        "period_type": period_type,
        "period_label": period_label,
        "claims": claims,
        "exposure_months": exposure_months,
        "rate_per_100000_member_months": claims / exposure_months * rate_per,
        "poisson_ci_lower_per_100000": ci_lower_count / exposure_months * rate_per,
        "poisson_ci_upper_per_100000": ci_upper_count / exposure_months * rate_per,
    }


def utilization_rates_by_period(ts, code, rate_per=100000):
    work = ts.copy()
    if "enrollment" not in work.columns:
        raise ValueError("utilization rates require an enrollment/exposure column.")

    work["month"] = pd.to_datetime(work["month"]).dt.to_period("M").dt.to_timestamp()
    work["y"] = pd.to_numeric(work["y"], errors="coerce").fillna(0)
    work["enrollment"] = pd.to_numeric(work["enrollment"], errors="coerce")
    work = work.dropna(subset=["month", "enrollment"])
    work = work[work["enrollment"] > 0].copy()

    rows = []
    period_specs = [
        ("month", work["month"].dt.to_period("M"), lambda period: period.strftime("%Y-%m")),
        ("quarter", work["month"].dt.to_period("Q"), lambda period: f"{period.year}-Q{period.quarter}"),
        ("year", work["month"].dt.to_period("Y"), lambda period: str(period.year)),
    ]

    for period_type, periods, label_func in period_specs:
        period_work = work.assign(_period=periods)
        grouped = period_work.groupby("_period", sort=True).agg(
            claims=("y", "sum"),
            exposure_months=("enrollment", "sum"),
        )
        for period, row in grouped.iterrows():
            rows.append(
                _period_rate_row(
                    code,
                    period_type,
                    label_func(period),
                    row["claims"],
                    row["exposure_months"],
                    rate_per=rate_per,
                )
            )

    return pd.DataFrame(rows, columns=UTILIZATION_RATE_COLUMNS)


def clean_summary_text(summary_text):
    text = str(summary_text).replace(chr(13) + chr(10), chr(10))
    marker = chr(10) + "Notes:" + chr(10)
    if marker in text:
        text = text.split(marker, 1)[0].rstrip()
    return text


def pick_font_size(summary_text, box_width_pts, box_height_pts):
    lines = summary_text.splitlines() or [""]
    max_len = max(len(line) for line in lines)
    line_count = len(lines)
    for size in (12, 11, 10, 9, 8, 7, 6):
        max_chars = box_width_pts / (0.60 * size)
        max_lines = box_height_pts / (1.20 * size)
        if max_len <= max_chars and line_count <= max_lines:
            return size
    return 6


def _month_to_quarter_label(month_str):
    ts = pd.Timestamp(month_str)
    q = (ts.month - 1) // 3 + 1
    return f"{ts.year}-Q{q}"


def plot_quarterly_rates_bars(rates, title, rate_per=100000, intervention_quarter=None):
    plot_df = rates[rates["period_type"].eq("quarter")].copy() if "period_type" in rates.columns else rates.copy()
    for col in ["rate_per_100000_member_months", "poisson_ci_lower_per_100000", "poisson_ci_upper_per_100000"]:
        plot_df[col] = pd.to_numeric(plot_df[col], errors="coerce")
    plot_df = plot_df.dropna(subset=["rate_per_100000_member_months"]).reset_index(drop=True)

    x = np.arange(len(plot_df))
    y = plot_df["rate_per_100000_member_months"].to_numpy(dtype=float)
    lower = np.maximum(y - plot_df["poisson_ci_lower_per_100000"].to_numpy(dtype=float), 0)
    upper = np.maximum(plot_df["poisson_ci_upper_per_100000"].to_numpy(dtype=float) - y, 0)

    fig_width = max(11, min(18, 0.36 * max(len(plot_df), 1)))
    fig, ax = plt.subplots(figsize=(fig_width, 6.1))
    ax.bar(
        x,
        y,
        color="#4C78A8",
        edgecolor="#26384D",
        linewidth=0.4,
        yerr=np.vstack([lower, upper]),
        error_kw={"ecolor": "#222222", "elinewidth": 0.9, "capsize": 2.5, "capthick": 0.9},
    )
    ax.set_title(title)
    ax.set_xlabel("Quarter")
    ax.set_ylabel(f"Claims per {rate_per:,} member-months")
    ax.set_xticks(x)
    ax.set_xticklabels(plot_df["period_label"], rotation=65, ha="right")
    ax.ticklabel_format(axis="y", style="plain", useOffset=False)
    ax.grid(axis="y", alpha=0.25)
    ax.set_axisbelow(True)

    if intervention_quarter and intervention_quarter in set(plot_df["period_label"]):
        intervention_pos = plot_df.index[plot_df["period_label"].eq(intervention_quarter)][0]
        ax.axvline(intervention_pos - 0.5, color="#C44E52", linestyle=":", linewidth=1.6)

    y_max = np.nanmax(np.r_[y + upper, [0.0]])
    ax.set_ylim(0, y_max * 1.10 if y_max > 0 else 1.0)
    fig.tight_layout()
    return fig


def add_quarterly_rates_slide(prs, target_title, rates, rate_per=100000, intervention_quarter=None, agg_codes=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide_title_text = f"Quarterly rate per {rate_per:,} member-months: {target_title}"

    title_box = slide.shapes.add_textbox(Inches(0.35), Inches(0.22), prs.slide_width - Inches(0.70), Inches(0.58))
    tf = title_box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = slide_title_text
    p.font.name = "Aptos"
    p.font.bold = True
    p.font.size = Pt(22)

    fig = plot_quarterly_rates_bars(rates, slide_title_text, rate_per=rate_per, intervention_quarter=intervention_quarter)
    image = io.BytesIO()
    fig.savefig(image, format="png", dpi=220, bbox_inches="tight")
    plt.close(fig)
    image.seek(0)
    slide.shapes.add_picture(image, Inches(0.42), Inches(0.95), prs.slide_width - Inches(0.84), Inches(5.15))

    note_box = slide.shapes.add_textbox(Inches(0.48), Inches(6.20), prs.slide_width - Inches(0.96), Inches(0.28))
    nf = note_box.text_frame
    nf.clear()
    note_parts = ["Error bars show exact 95% Poisson confidence intervals."]
    if agg_codes:
        note_parts.append("Codes in aggregation: " + ", ".join(str(code) for code in agg_codes))
    nf.paragraphs[0].text = " ".join(note_parts)
    nf.paragraphs[0].font.name = "Aptos"
    nf.paragraphs[0].font.size = Pt(8)

    return prs.slides._sldIdLst[-1]


def plot_poisson(ts, title, intervention_month, rate_per=100000):
    return plot_normalized_its(
        ts,
        title,
        intervention_month,
        fitted_label=f"Fitted Poisson ITS rate per {rate_per:,} member-months",
        rate_per=rate_per,
    )


def plot_negative_binomial(ts, title, intervention_month, rate_per=100000):
    return plot_normalized_its(
        ts,
        title,
        intervention_month,
        fitted_label=f"Fitted Negative Binomial ITS rate per {rate_per:,} member-months",
        rate_per=rate_per,
    )


def plot_normalized_its(ts, title, intervention_month, fitted_label, rate_per=100000):
    fig, ax = plt.subplots(figsize=(10, 8))
    ax.plot(
        ts["month"],
        ts["rate"],
        marker="o",
        linewidth=1.5,
        label=f"Observed claims per {rate_per:,} member-months",
    )
    ax.plot(
        ts["month"],
        ts["rate_hat"],
        linestyle="--",
        linewidth=2,
        label=fitted_label,
    )
    ax.axvline(pd.Timestamp(intervention_month), color="red", linestyle=":", linewidth=2, label="Intervention")
    ax.set_xlabel("Month")
    ax.set_ylabel(f"Claims per {rate_per:,} Medicaid member-months")
    y_max = float(np.nanmax(np.r_[ts["rate"].to_numpy(dtype=float), ts["rate_hat"].to_numpy(dtype=float), [0.0]]))
    ax.set_ylim(0, y_max * 1.05 if y_max > 0 else 1.0)
    ax.ticklabel_format(axis="y", style="plain", useOffset=False)
    ax.yaxis.set_major_formatter(FuncFormatter(lambda x, _: f"{x:,.2f}".rstrip("0").rstrip(".")))
    ax.set_title(title)
    ax.legend()
    ax.grid(alpha=0.25)
    plt.xticks(rotation=45)
    plt.tight_layout()
    return fig


def result_n_obs_lookup(results_path):
    results = pd.read_csv(results_path, dtype={"HCPCS_CODE": str})
    lookup = {}
    for _, row in results.iterrows():
        code = str(row.get("HCPCS_CODE", "")).strip().upper()
        n_obs = pd.to_numeric(row.get("N_obs"), errors="coerce")
        if code and pd.notna(n_obs):
            lookup[code] = int(n_obs)
    return lookup


def choose_series(claims_full, claims_drop, code_col, result_lookup):
    full = build_ts(claims_full, code_col)
    dropped = build_ts(claims_drop, code_col)
    target_n = result_lookup.get(str(code_col).strip().upper())

    if target_n is not None:
        if int(dropped["y"].notna().sum()) == target_n:
            return dropped
        if int(full["y"].notna().sum()) == target_n:
            return full

    return dropped if len(dropped) else full


def is_aggregate_code(code_col):
    return str(code_col).strip().upper().startswith("AGG")


def aggregation_codes(code_col):
    target = str(code_col).strip().upper()
    for agg_name, codes in AGG_MAP.items():
        if agg_name.strip().upper() == target:
            return codes
    return []


def safe_exp(value):
    if pd.isna(value) or value < -700 or value > 700:
        return np.nan
    return float(np.exp(value))


def format_numeric(value, digits=3):
    if pd.isna(value) or not np.isfinite(value):
        return ""
    return f"{value:.{digits}g}"


def negative_binomial_rows(model):
    rows = []
    for term in ["const", "T", "X_t", "post_time", "alpha"]:
        if term not in model.params.index:
            continue
        coef = float(model.params[term])
        se = float(model.bse[term]) if term in model.bse.index else np.nan
        pval = float(model.pvalues[term]) if term in model.pvalues.index else np.nan
        irr = safe_exp(coef) if term != "alpha" else np.nan
        rows.append(
            [
                term,
                format_numeric(coef),
                format_numeric(se),
                format_numeric(pval),
                format_numeric(irr),
            ]
        )
    return rows


def set_cell_text(cell, text, font_size=7, bold=False):
    cell.text = str(text)
    for para in cell.text_frame.paragraphs:
        para.font.name = "Aptos"
        para.font.size = Pt(font_size)
        para.font.bold = bold


def add_negative_binomial_table(slide, geom, nb_model, top, height):
    title_box = slide.shapes.add_textbox(
        int(geom["summary_left"]),
        int(top),
        int(geom["summary_width"]),
        int(170000),
    )
    tf = title_box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = "Negative Binomial ITS"
    p.font.bold = True
    p.font.size = Pt(10)

    rows = negative_binomial_rows(nb_model)
    table_shape = slide.shapes.add_table(
        len(rows) + 1,
        5,
        int(geom["summary_left"]),
        int(top + 190000),
        int(geom["summary_width"]),
        int(height - 190000),
    )
    table = table_shape.table
    headers = ["Term", "Coef", "SE", "P>|z|", "IRR"]
    for col_idx, header in enumerate(headers):
        set_cell_text(table.cell(0, col_idx), header, font_size=7, bold=True)
    for row_idx, row in enumerate(rows, start=1):
        for col_idx, value in enumerate(row):
            set_cell_text(table.cell(row_idx, col_idx), value, font_size=7)


def add_aggregate_codes_box(slide, geom, codes, top, height):
    if not codes:
        return
    box = slide.shapes.add_textbox(
        int(geom["summary_left"]),
        int(top),
        int(geom["summary_width"]),
        int(height),
    )
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = "Codes in aggregation: " + ", ".join(codes)
    p.font.name = "Aptos"
    p.font.size = Pt(7)


def add_negative_binomial_error(slide, geom, error_text, top, height):
    box = slide.shapes.add_textbox(
        int(geom["summary_left"]),
        int(top),
        int(geom["summary_width"]),
        int(height),
    )
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = "Negative Binomial ITS could not be fit: " + str(error_text)
    p.font.name = "Aptos"
    p.font.size = Pt(8)
    p.font.bold = True


def add_poisson_slide(
    prs,
    target,
    ts,
    model,
    intervention_month,
    rate_per=100000,
    agg_codes=None,
):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    geom = target["template"]
    poisson_title = f"Normalized Poisson ITS (per {rate_per:,} member-months): {target['title']}"

    title_box = slide.shapes.add_textbox(
        int(geom["title_left"]),
        int(geom["title_top"]),
        int(geom["title_width"]),
        int(geom["title_height"]),
    )
    tf = title_box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = poisson_title
    p.font.bold = True
    p.font.size = Pt(24)

    fig = plot_poisson(ts, poisson_title, intervention_month, rate_per=rate_per)
    image = io.BytesIO()
    fig.savefig(image, format="png", dpi=200, bbox_inches="tight")
    plt.close(fig)
    image.seek(0)
    slide.shapes.add_picture(
        image,
        int(geom["plot_left"]),
        int(geom["plot_top"]),
        int(geom["plot_width"]),
        int(geom["plot_height"]),
    )

    is_aggregate = bool(agg_codes)
    summary_height = int(geom["summary_height"])
    if is_aggregate:
        bottom_margin = 220000
        codes_height = 520000
        codes_top = int(prs.slide_height) - bottom_margin - codes_height
        summary_height = min(summary_height, max(900000, codes_top - int(geom["summary_top"]) - 120000))

    summary_box = slide.shapes.add_textbox(
        int(geom["summary_left"]),
        int(geom["summary_top"]),
        int(geom["summary_width"]),
        summary_height,
    )
    sf = summary_box.text_frame
    sf.clear()
    sf.word_wrap = False
    sf.margin_left = 0
    sf.margin_right = 0
    sf.margin_top = 0
    sf.margin_bottom = 0
    summary_text = clean_summary_text(model.summary().as_text())
    sf.text = summary_text
    font_size = pick_font_size(
        summary_text,
        int(geom["summary_width"]) / EMU_PER_PT,
        summary_height / EMU_PER_PT,
    )
    for para in sf.paragraphs:
        para.font.name = "Consolas"
        para.font.size = Pt(font_size)

    if is_aggregate:
        add_aggregate_codes_box(slide, geom, agg_codes, codes_top, codes_height)

    return prs.slides._sldIdLst[-1]


def add_negative_binomial_slide(
    prs,
    target,
    ts,
    model,
    intervention_month,
    rate_per=100000,
    agg_codes=None,
    nb_error=None,
):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    geom = target["template"]
    nb_title = f"Normalized Negative Binomial ITS (per {rate_per:,} member-months): {target['title']}"

    title_box = slide.shapes.add_textbox(
        int(geom["title_left"]),
        int(geom["title_top"]),
        int(geom["title_width"]),
        int(geom["title_height"]),
    )
    tf = title_box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = nb_title
    p.font.bold = True
    p.font.size = Pt(24)

    if model is not None:
        fig = plot_negative_binomial(ts, nb_title, intervention_month, rate_per=rate_per)
        image = io.BytesIO()
        fig.savefig(image, format="png", dpi=200, bbox_inches="tight")
        plt.close(fig)
        image.seek(0)
        slide.shapes.add_picture(
            image,
            int(geom["plot_left"]),
            int(geom["plot_top"]),
            int(geom["plot_width"]),
            int(geom["plot_height"]),
        )

    is_aggregate = bool(agg_codes)
    summary_height = int(geom["summary_height"])
    if is_aggregate:
        bottom_margin = 220000
        codes_height = 520000
        codes_top = int(prs.slide_height) - bottom_margin - codes_height
        summary_height = min(summary_height, max(900000, codes_top - int(geom["summary_top"]) - 120000))

    summary_box = slide.shapes.add_textbox(
        int(geom["summary_left"]),
        int(geom["summary_top"]),
        int(geom["summary_width"]),
        summary_height,
    )
    sf = summary_box.text_frame
    sf.clear()
    sf.word_wrap = False
    sf.margin_left = 0
    sf.margin_right = 0
    sf.margin_top = 0
    sf.margin_bottom = 0
    if model is not None:
        summary_text = clean_summary_text(model.summary().as_text())
    else:
        summary_text = "Negative Binomial ITS could not be fit:\n" + str(nb_error)
    sf.text = summary_text
    font_size = pick_font_size(
        summary_text,
        int(geom["summary_width"]) / EMU_PER_PT,
        summary_height / EMU_PER_PT,
    )
    for para in sf.paragraphs:
        para.font.name = "Consolas"
        para.font.size = Pt(font_size)

    if is_aggregate:
        add_aggregate_codes_box(slide, geom, agg_codes, codes_top, codes_height)

    return prs.slides._sldIdLst[-1]


def update_build_summary(prs, summary_text):
    for slide in prs.slides:
        if slide_title(slide).strip().lower() != "build summary":
            continue
        text_shapes = [shape for shape in slide.shapes if getattr(shape, "has_text_frame", False)]
        text_shapes = sorted(text_shapes, key=lambda shape: (int(shape.top), int(shape.left)))
        if text_shapes:
            text_shapes[0].text_frame.text = "Build Summary"
        if len(text_shapes) > 1:
            body = max(text_shapes[1:], key=lambda shape: int(shape.width) * int(shape.height))
            body.text_frame.text = summary_text
        return


def build_poisson_deck(
    source_pptx="its_all_codes_reportRetsefV7_with_normalized.pptx",
    output_pptx="its_all_codes_reportRetsefV7_normalized_poisson.pptx",
    utilization_rates_path=None,
    claims_path="claims_pivot.csv",
    enrollment_path="Enrollees_Merged.xlsx",
    poisson_results_path="its_results_poisson.csv",
    intervention_month="2021-03-01",
    rate_per=100000,
    drop_months_for_robust=("2020-03-01", "2020-05-31"),
):
    source_pptx = Path(source_pptx)
    output_pptx = Path(output_pptx)
    if utilization_rates_path is None:
        utilization_rates_path = output_pptx.with_name(f"{output_pptx.stem}_utilization_rates.csv")
    utilization_rates_path = Path(utilization_rates_path)
    claims_full = load_claims_pivot(claims_path)
    enrollment = load_enrollment(enrollment_path)
    claims_drop = claims_full[
        ~claims_full.index.isin(
            pd.date_range(pd.Timestamp(drop_months_for_robust[0]), pd.Timestamp(drop_months_for_robust[1]), freq="D")
        )
    ].copy()
    result_lookup = result_n_obs_lookup(poisson_results_path)

    intervention_quarter = _month_to_quarter_label(intervention_month)

    prs = Presentation(str(source_pptx))
    original_sld_ids = list(prs.slides._sldIdLst)
    agg_lookup = aggregate_title_lookup(claims_full)

    targets = []
    seen = set()
    for idx, slide in enumerate(prs.slides):
        if slide._element.get("show") == "0":
            continue
        title = slide_title(slide)
        code_col = slide_code(title, claims_full, agg_lookup)
        if code_col is None or code_col in seen:
            continue
        seen.add(code_col)
        targets.append(
            {
                "orig_index": idx,
                "title": title,
                "code_col": code_col,
                "template": extract_template_geometry(slide),
            }
        )

    created_ids = {}
    utilization_rate_frames = []
    skipped = []
    nb_failed = []
    for target in targets:
        code_col = target["code_col"]
        try:
            ts = choose_series(claims_full, claims_drop, code_col, result_lookup)
            if ts.empty:
                raise ValueError("no observations after trimming")
            fitted_ts, model = fit_normalized_poisson(ts, enrollment, intervention_month, rate_per=rate_per)
            target_rates = utilization_rates_by_period(fitted_ts, code_col, rate_per=rate_per)
            utilization_rate_frames.append(target_rates)
            agg_codes = aggregation_codes(code_col) if is_aggregate_code(code_col) else []
            nb_model = None
            nb_ts = None
            nb_error = None
            try:
                nb_ts, nb_model = fit_normalized_negative_binomial(
                    fitted_ts,
                    intervention_month,
                    rate_per=rate_per,
                )
            except Exception as exc:
                nb_error = exc
                nb_failed.append((target["title"], str(exc)))

            poisson_id = add_poisson_slide(
                prs,
                target,
                fitted_ts,
                model,
                intervention_month,
                rate_per=rate_per,
                agg_codes=agg_codes,
            )
            nb_id = add_negative_binomial_slide(
                prs,
                target,
                nb_ts if nb_ts is not None else fitted_ts,
                nb_model,
                intervention_month,
                rate_per=rate_per,
                agg_codes=agg_codes,
                nb_error=nb_error,
            )
            rates_id = add_quarterly_rates_slide(
                prs,
                target["title"],
                target_rates,
                rate_per=rate_per,
                intervention_quarter=intervention_quarter,
                agg_codes=agg_codes if agg_codes else None,
            )
            created_ids[code_col] = [poisson_id, nb_id, rates_id]
        except Exception as exc:
            skipped.append((target["title"], str(exc)))

    desired_order = []
    seen_order = set()
    kept_transitions = 0
    removed_data_or_normalized = 0
    removed_hidden = 0

    for idx, slide in enumerate(prs.slides):
        if idx >= len(original_sld_ids):
            continue

        if slide._element.get("show") == "0":
            removed_hidden += 1
            continue

        title = slide_title(slide)
        code_col = slide_code(title, claims_full, agg_lookup)
        if title.lower().startswith("normalized its"):
            removed_data_or_normalized += 1
            continue

        if code_col is not None:
            removed_data_or_normalized += 1
            if code_col in created_ids and code_col not in seen_order:
                desired_order.extend(created_ids[code_col])
                seen_order.add(code_col)
            continue

        desired_order.append(original_sld_ids[idx])
        kept_transitions += 1

    sld_id_list = prs.slides._sldIdLst
    for sld_id in list(sld_id_list):
        sld_id_list.remove(sld_id)
    for sld_id in desired_order:
        sld_id_list.append(sld_id)

    summary_text = (
        f"Output: {output_pptx.resolve()}\n"
        f"Utilization rates: {utilization_rates_path.resolve()}\n"
        f"Source style/template: {source_pptx.name}\n"
        "Models: Poisson GLM and Negative Binomial interrupted time series with enrollment exposure\n"
        f"Display scale: claims per {rate_per:,} Medicaid member-months\n"
        f"Intervention month: {intervention_month}\n"
        f"Normalized Poisson slides created: {len(created_ids)}\n"
        f"Normalized Negative Binomial slides created: {len(created_ids)}\n"
        f"Negative Binomial fit failures shown as error slides: {len(nb_failed)}\n"
        f"Transition/title slides kept: {kept_transitions}\n"
        f"Original data/normalized slides removed: {removed_data_or_normalized}\n"
        f"Hidden slides removed: {removed_hidden}\n"
        f"Skipped targets: {len(skipped)}"
    )
    update_build_summary(prs, summary_text)

    output_pptx.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_pptx))
    utilization_rates = (
        pd.concat(utilization_rate_frames, ignore_index=True)
        if utilization_rate_frames
        else pd.DataFrame(columns=UTILIZATION_RATE_COLUMNS)
    )
    utilization_rates_path.parent.mkdir(parents=True, exist_ok=True)
    utilization_rates.to_csv(utilization_rates_path, index=False)
    return {
        "source_pptx": str(source_pptx),
        "output_pptx": str(output_pptx),
        "utilization_rates_path": str(utilization_rates_path),
        "utilization_rate_rows": int(len(utilization_rates)),
        "target_visible_data_groups": len(targets),
        "normalized_poisson_slides_created": len(created_ids),
        "normalized_negative_binomial_slides_created": len(created_ids),
        "negative_binomial_fit_failures": nb_failed,
        "transition_slides_kept": kept_transitions,
        "removed_data_or_normalized_slides": removed_data_or_normalized,
        "removed_hidden_slides": removed_hidden,
        "skipped": skipped,
    }


if __name__ == "__main__":
    report = build_poisson_deck()
    print(report)
