"""
Remake all analysis + graphs in
its_all_codes_reportRetsefV12_with_quarterly_and_yearly_rates.pptx
using data truncated at Dec 2023 (claims_pivot.csv now ends 2023-12).

Methodology reverse-engineered from V12 and applied uniformly:
  * intervention month  = 2021-03-01 (quarterly bar line at 2021-Q1)
  * drop-2020 robustness: exclude 2020-03/04/05 from every series (ITS + rates)
  * Poisson GLM ITS + Negative Binomial ITS, enrollment as exposure
  * quarterly + yearly exact-Poisson utilization rates

For each analysis slide we rebuild its own series from the current data and swap
the plot image + regression summary (ITS) or the bar chart + yearly table
(quarterly) in place, preserving every slide's geometry, order and wording.
"""

import io
import re
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.ticker import FuncFormatter
import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt

import build_poisson_pptx as P
import add_yearly_rates as A

SOURCE_DECK = "its_all_codes_reportRetsefV12_with_quarterly_and_yearly_rates.pptx"
OUTPUT_DECK = "its_all_codes_reportRetsefV13_through_2023.pptx"
CLAIMS_PATH = "claims_pivot.csv"
ENROLL_PATH = "Enrollees_Merged.xlsx"
INTERVENTION = "2021-03-01"
INTERVENTION_QUARTER = "2021-Q1"
DROP = ("2020-03-01", "2020-05-31")
CUTOFF = pd.Timestamp("2023-12-01")
RATE_PER = 100_000

def trim_trailing_zeros(ts):
    """Cut a run of trailing zero/NaN months so a discontinued code ends at its
    last real observation (matches the original ITS series construction). Fixes
    the spike a log-link ITS produces when chasing a long tail of zeros."""
    ts = ts.reset_index(drop=True)
    y = pd.to_numeric(ts["y"], errors="coerce").fillna(0).to_numpy()
    nz = np.nonzero(y > 0)[0]
    if len(nz) == 0:
        return ts.iloc[0:0].copy()
    return ts.loc[: nz[-1]].reset_index(drop=True)


POIS_PREFIX = re.compile(r"^Normalized Poisson ITS \([^)]*\):\s*")
NB_PREFIX = re.compile(r"^Normalized Negative Binomial ITS \([^)]*\):\s*")
PICTURE = 13


# ---------- plotting (ported verbatim in style from notebook cells 18/19/22) ----------

def _fmt_yaxis(ax):
    ax.ticklabel_format(axis="y", style="plain", useOffset=False)
    ax.yaxis.set_major_formatter(FuncFormatter(lambda x, _: f"{x:,.2f}".rstrip("0").rstrip(".")))


def plot_its(ts, title, fitted_label):
    fig, ax = plt.subplots(figsize=(10, 8))
    ax.plot(ts["month"], ts["rate"], marker="o", linewidth=1.5,
            label=f"Observed claims per {RATE_PER:,} member-months")
    ax.plot(ts["month"], ts["rate_hat"], linestyle="--", linewidth=2, label=fitted_label)
    ax.axvline(pd.Timestamp(INTERVENTION), color="red", linestyle=":", linewidth=2, label="Intervention")
    ax.set_xlabel("Month")
    ax.set_ylabel(f"Claims per {RATE_PER:,} Medicaid member-months")
    ymax = float(np.nanmax(np.r_[ts["rate"].to_numpy(float), ts["rate_hat"].to_numpy(float), [0.0]]))
    ax.set_ylim(0, ymax * 1.05 if ymax > 0 else 1.0)
    _fmt_yaxis(ax)
    ax.set_title(title)
    ax.legend()
    ax.grid(alpha=0.25)
    plt.xticks(rotation=45)
    plt.tight_layout()
    return fig


def plot_quarterly_bars(rates, title):
    plot_df = rates[rates["period_type"].eq("quarter")].copy()
    for col in ["rate_per_100000_member_months", "poisson_ci_lower_per_100000", "poisson_ci_upper_per_100000"]:
        plot_df[col] = pd.to_numeric(plot_df[col], errors="coerce")
    plot_df = plot_df.dropna(subset=["rate_per_100000_member_months"]).reset_index(drop=True)
    x = np.arange(len(plot_df))
    y = plot_df["rate_per_100000_member_months"].to_numpy(float)
    lower = np.maximum(y - plot_df["poisson_ci_lower_per_100000"].to_numpy(float), 0)
    upper = np.maximum(plot_df["poisson_ci_upper_per_100000"].to_numpy(float) - y, 0)
    fig_width = max(11, min(18, 0.36 * max(len(plot_df), 1)))
    fig, ax = plt.subplots(figsize=(fig_width, 6.1))
    ax.bar(x, y, color="#4C78A8", edgecolor="#26384D", linewidth=0.4,
           yerr=np.vstack([lower, upper]),
           error_kw={"ecolor": "#222222", "elinewidth": 0.9, "capsize": 2.5, "capthick": 0.9})
    ax.set_title(title)
    ax.set_xlabel("Quarter")
    ax.set_ylabel(f"Claims per {RATE_PER:,} member-months")
    ax.set_xticks(x)
    ax.set_xticklabels(plot_df["period_label"], rotation=65, ha="right")
    ax.ticklabel_format(axis="y", style="plain", useOffset=False)
    ax.grid(axis="y", alpha=0.25)
    ax.set_axisbelow(True)
    if INTERVENTION_QUARTER in set(plot_df["period_label"]):
        pos = plot_df.index[plot_df["period_label"].eq(INTERVENTION_QUARTER)][0]
        ax.axvline(pos - 0.5, color="#C44E52", linestyle=":", linewidth=1.6)
    y_max = np.nanmax(np.r_[y + upper, [0.0]])
    ax.set_ylim(0, y_max * 1.10 if y_max > 0 else 1.0)
    fig.tight_layout()
    return fig


def fig_to_png(fig, dpi):
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=dpi, bbox_inches="tight")
    plt.close(fig)
    buf.seek(0)
    return buf


# ---------- in-place shape helpers ----------

def find_picture(slide):
    for sh in slide.shapes:
        if sh.shape_type == PICTURE:
            return sh
    return None


def replace_picture(slide, png, geom):
    old = find_picture(slide)
    if old is not None:
        old._element.getparent().remove(old._element)
    slide.shapes.add_picture(png, geom[0], geom[1], geom[2], geom[3])


def set_summary_text(box, text):
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    text = P.clean_summary_text(text)
    tf.text = text
    size = P.pick_font_size(text, box.width / 12700, box.height / 12700)
    for para in tf.paragraphs:
        para.font.name = "Consolas"
        para.font.size = Pt(size)


def update_title_hint(box, new_n):
    """Rewrite '(N months of data)' -> new count on whichever paragraph holds it,
    preserving that paragraph's font. Handles multi-line (wrapped) titles."""
    for para in box.text_frame.paragraphs:
        runs = para.runs
        full = "".join(r.text for r in runs)
        if "months of data" not in full:
            continue
        new_full = re.sub(r"\(\d+\s+months of data\)", f"({new_n} months of data)", full)
        if new_full == full or not runs:
            continue
        f = runs[0].font
        name, size, bold = f.name, f.size, f.bold
        color = None
        try:
            if f.color and f.color.type is not None:
                color = f.color.rgb
        except Exception:
            pass
        runs[0].text = new_full
        for extra in runs[1:]:
            extra.text = ""
        runs[0].font.name = name
        if size is not None:
            runs[0].font.size = size
        runs[0].font.bold = bold
        if color is not None:
            runs[0].font.color.rgb = color


def title_shape(slide, prefix_re=None, startswith=None):
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        t = A._single_line(sh.text_frame.text)
        if prefix_re and prefix_re.match(t):
            return sh, t
        if startswith and t.startswith(startswith):
            return sh, t
    return None, None


# ---------- yearly table (same layout as add_yearly_rates) ----------

def remove_old_yearly(slide):
    for sh in list(slide.shapes):
        if getattr(sh, "has_table", False) and sh.has_table:
            sh._element.getparent().remove(sh._element)
        elif sh.has_text_frame and A._single_line(sh.text_frame.text).startswith("Yearly rate per"):
            sh._element.getparent().remove(sh._element)


def add_yearly_table(slide, prs, ydf):
    from pptx.dml.color import RGBColor
    HEAD = (0x26, 0x38, 0x4D); POST = (0xF2, 0xDE, 0xDE); PRE = (0xEE, 0xF2, 0xF7)
    ydf = ydf[pd.to_numeric(ydf["exposure_months"], errors="coerce") > 0]
    years = ydf["period_label"].tolist()
    n = len(years)
    if n == 0:
        return
    cap = slide.shapes.add_textbox(Inches(0.48), Inches(6.52), prs.slide_width - Inches(0.96), Inches(0.22))
    cp = cap.text_frame.paragraphs[0]
    cp.text = "Yearly rate per 100,000 member-months (exact 95% Poisson CIs):"
    cp.font.name = "Aptos"; cp.font.bold = True; cp.font.size = Pt(9)
    tbl_w = prs.slide_width - Inches(0.96)
    tbl = slide.shapes.add_table(3, n + 1, Inches(0.48), Inches(6.74), tbl_w, Inches(0.70)).table
    tbl.first_row = False; tbl.horz_banding = False
    label_w = Inches(1.30); year_w = int((tbl_w - label_w) / n)
    tbl.columns[0].width = label_w
    for c in range(1, n + 1):
        tbl.columns[c].width = year_w
    for r in range(3):
        tbl.rows[r].height = Inches(0.70 / 3)
    A._set_cell(tbl.cell(0, 0), "Year", bold=True, fill=HEAD, align="left")
    A._set_cell(tbl.cell(1, 0), "Rate", bold=True, align="left")
    A._set_cell(tbl.cell(2, 0), "95% CI", bold=True, align="left")
    tbl.cell(0, 0).text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for j, (_, row) in enumerate(ydf.iterrows(), start=1):
        yr = int(row["period_label"]); post = yr >= 2021
        A._set_cell(tbl.cell(0, j), str(yr), bold=True, fill=(POST if post else PRE))
        A._set_cell(tbl.cell(1, j), A._fmt_rate(row["rate_per_100000_member_months"]))
        ci = f"{A._fmt_rate(row['poisson_ci_lower_per_100000'])}–{A._fmt_rate(row['poisson_ci_upper_per_100000'])}"
        A._set_cell(tbl.cell(2, j), ci, size=7.5)


# ---------- main ----------

def main():
    claims_full = P.load_claims_pivot(CLAIMS_PATH)
    claims_full = claims_full[claims_full.index <= CUTOFF]
    enrollment = P.load_enrollment(ENROLL_PATH)
    claims_drop = claims_full[
        ~claims_full.index.isin(pd.date_range(pd.Timestamp(DROP[0]), pd.Timestamp(DROP[1]), freq="D"))
    ].copy()
    agg_lookup = P.aggregate_title_lookup(claims_full)

    def resolve_col(slide, idx, base):
        col, tt, _ = A._claim_column_from_base(slide, idx, base, claims_full, agg_lookup)
        if col not in claims_drop.columns and col in claims_full.columns:
            claims_drop[col] = claims_full.loc[claims_drop.index, col]
        return col

    def series(col):
        return P.build_ts(claims_drop, col)

    prs = Presentation(SOURCE_DECK)
    report = {"poisson": 0, "nb": 0, "nb_fail": 0, "quarterly": 0, "skipped": []}

    for idx, slide in enumerate(prs.slides):
        tsh, ttext = title_shape(slide, prefix_re=POIS_PREFIX)
        if tsh is not None:
            base = POIS_PREFIX.sub("", ttext)
            try:
                col = resolve_col(slide, idx + 1, base)
                ts = trim_trailing_zeros(series(col))
                fitted, model = P.fit_normalized_poisson(ts, enrollment, INTERVENTION)
                new_title = re.sub(r"\(\d+\s+months of data\)", f"({len(fitted)} months of data)", ttext)
                geom = (find_picture(slide).left, find_picture(slide).top,
                        find_picture(slide).width, find_picture(slide).height)
                fig = plot_its(fitted, new_title, f"Fitted Poisson ITS rate per {RATE_PER:,} member-months")
                replace_picture(slide, fig_to_png(fig, 200), geom)
                sumbox = next(sh for sh in slide.shapes
                              if sh.has_text_frame and "Regression Results" in sh.text_frame.text)
                set_summary_text(sumbox, model.summary().as_text())
                update_title_hint(tsh, len(fitted))
                report["poisson"] += 1
            except Exception as exc:
                report["skipped"].append(("POIS", idx, base[:40], str(exc)))
            continue

        tsh, ttext = title_shape(slide, prefix_re=NB_PREFIX)
        if tsh is not None:
            base = NB_PREFIX.sub("", ttext)
            try:
                col = resolve_col(slide, idx + 1, base)
                ts = trim_trailing_zeros(series(col))
                fitted_p, _ = P.fit_normalized_poisson(ts, enrollment, INTERVENTION)
                nb_ts, nb_model = P.fit_normalized_negative_binomial(fitted_p, INTERVENTION)
                new_title = re.sub(r"\(\d+\s+months of data\)", f"({len(nb_ts)} months of data)", ttext)
                geom = (find_picture(slide).left, find_picture(slide).top,
                        find_picture(slide).width, find_picture(slide).height)
                fig = plot_its(nb_ts, new_title, f"Fitted negative binomial ITS rate per {RATE_PER:,} member-months")
                replace_picture(slide, fig_to_png(fig, 200), geom)
                sumbox = next(sh for sh in slide.shapes
                              if sh.has_text_frame and "Regression Results" in sh.text_frame.text)
                set_summary_text(sumbox, nb_model.summary().as_text())
                update_title_hint(tsh, len(nb_ts))
                report["nb"] += 1
            except Exception as exc:
                report["nb_fail"] += 1
                report["skipped"].append(("NB", idx, base[:40], str(exc)))
            continue

        tsh, ttext = title_shape(slide, startswith=A.QUARTERLY_PREFIX)
        if tsh is not None:
            base = ttext[len(A.QUARTERLY_PREFIX):].strip()
            try:
                col = resolve_col(slide, idx + 1, base)
                ts = series(col)
                work = ts.copy()
                work["month"] = pd.to_datetime(work["month"]).dt.to_period("M").dt.to_timestamp()
                enr = enrollment.copy()
                enr["month"] = pd.to_datetime(enr["month"]).dt.to_period("M").dt.to_timestamp()
                work = work.merge(enr[["month", "enrollment"]], on="month", how="left", validate="many_to_one")
                rates = P.utilization_rates_by_period(work, base, rate_per=RATE_PER)
                q = rates[rates["period_type"].eq("quarter")].reset_index(drop=True)
                y = rates[rates["period_type"].eq("year")].reset_index(drop=True)
                pic = find_picture(slide)
                geom = (pic.left, pic.top, pic.width, pic.height)
                new_title = re.sub(r"\(\d+\s+months of data\)", f"({int(ts['y'].notna().sum())} months of data)", ttext)
                fig = plot_quarterly_bars(q, new_title)
                remove_old_yearly(slide)
                replace_picture(slide, fig_to_png(fig, 220), geom)
                add_yearly_table(slide, prs, y)
                update_title_hint(tsh, int(ts["y"].notna().sum()))
                report["quarterly"] += 1
            except Exception as exc:
                report["skipped"].append(("QUART", idx, base[:40], str(exc)))
            continue

    prs.save(OUTPUT_DECK)
    return report


if __name__ == "__main__":
    rep = main()
    print("Poisson ITS slides remade   :", rep["poisson"])
    print("Neg-Binomial ITS slides remade:", rep["nb"], "(fit failures:", rep["nb_fail"], ")")
    print("Quarterly+yearly slides remade:", rep["quarterly"])
    print("skipped:", len(rep["skipped"]))
    for s in rep["skipped"]:
        print("   ", s)
    print("\nsaved ->", OUTPUT_DECK)
