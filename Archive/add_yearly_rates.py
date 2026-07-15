"""
Add yearly Poisson utilization rates to the existing quarterly-rate slides of the
V12 deck.

The quarterly slides were built (notebook cell 22) by reconstructing a per-code /
per-aggregation monthly time series and computing exact-Poisson rates per period.
We reuse the *identical* series-selection logic here so the yearly numbers are
guaranteed consistent with the quarterly bars already shown.  utilization_rates_by_period
returns month/quarter/year rows from the same underlying series, so the yearly rate for
a year is exactly sum(claims that year) / sum(exposure that year) * 100,000.
"""

import re
from collections import defaultdict
from pathlib import Path

import pandas as pd
from pptx import Presentation

import build_poisson_pptx as poisson_deck

SOURCE_PPTX = Path("its_all_codes_reportRetsefV12_with_quarterly_rates(1).pptx")
CLAIMS_PATH = Path("claims_pivot.csv")
ENROLLMENT_PATH = Path("Enrollees_Merged.xlsx")
RATE_PER = 100_000
DROP_MONTHS_FOR_ROBUST = ("2020-03-01", "2020-05-31")
QUARTERLY_PREFIX = f"Quarterly rate per {RATE_PER:,} member-months:"


# ---- helpers copied verbatim from notebook cell 22 -------------------------

def _single_line(text):
    return re.sub(r"\s+", " ", str(text)).strip()


def _quarterly_slide_title(slide):
    texts = []
    try:
        if slide.shapes.title is not None and slide.shapes.title.text and slide.shapes.title.text.strip():
            texts.append(slide.shapes.title.text.strip())
    except Exception:
        pass
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False) and shape.text and shape.text.strip():
            text = shape.text.strip()
            if text not in texts:
                texts.append(text)
    normalized_candidates = [
        text for text in texts
        if re.match(r"^\s*Normalized\s+(?:Poisson|Negative\s+Binomial)\s+ITS\b", text, flags=re.IGNORECASE)
    ]
    if normalized_candidates:
        return max(normalized_candidates, key=lambda text: len(_single_line(text)))
    return texts[0].splitlines()[0] if texts else ""


def _is_normalized_poisson_title(title):
    return bool(re.match(r"^\s*Normalized\s+Poisson\s+ITS\b", str(title), flags=re.IGNORECASE))


def _base_title_from_normalized_title(title):
    text = _single_line(title)
    match = re.match(
        r"^Normalized\s+(?:Poisson|Negative\s+Binomial)\s+ITS\s*\([^)]*\):\s*(.*)$",
        text,
        flags=re.IGNORECASE,
    )
    return match.group(1).strip() if match else text


def _codes_from_aggregation_box(slide):
    marker = "codes in aggregation:"
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        text = _single_line(shape.text)
        pos = text.lower().find(marker)
        if pos < 0:
            continue
        raw_codes = text[pos + len(marker):]
        return re.findall(r"\b[A-Z0-9]{4,5}\b", raw_codes.upper())
    return []


def _safe_column_name(label, idx):
    slug = re.sub(r"[^A-Za-z0-9]+", "_", str(label)).strip("_")
    return f"Quarterly_Custom_Agg_{idx:03d}_{slug[:60]}"


def _add_custom_aggregation_column(claims_df, label, codes, idx):
    code_col = _safe_column_name(label, idx)
    lookup = {str(col).strip().upper(): col for col in claims_df.columns}
    present = [lookup[str(code).strip().upper()] for code in codes if str(code).strip().upper() in lookup]
    claims_df[code_col] = claims_df[present].fillna(0).sum(axis=1) if present else 0.0
    return code_col, present


def _claim_column_from_base(slide, idx, base_title, claims_df, agg_lookup):
    agg_codes = _codes_from_aggregation_box(slide)

    hcpcs_match = re.search(r"\bHCPCS\s+([A-Z0-9]+)\b", base_title, flags=re.IGNORECASE)
    if hcpcs_match:
        key = hcpcs_match.group(1).strip().upper()
        lookup = {str(col).strip().upper(): col for col in claims_df.columns}
        if key not in lookup:
            raise ValueError(f"HCPCS {key} was not found in {CLAIMS_PATH}.")
        return lookup[key], base_title, []

    aggregate_key = poisson_deck.norm_text(base_title)
    if aggregate_key in agg_lookup:
        code_col = agg_lookup[aggregate_key]
        return code_col, base_title, agg_codes or poisson_deck.aggregation_codes(code_col)

    if agg_codes:
        code_col, present = _add_custom_aggregation_column(claims_df, base_title, agg_codes, idx)
        return code_col, base_title, [str(code) for code in present]

    raise ValueError(f"Could not identify a code or aggregation from slide title: {base_title!r}")


def _n_obs_hint(title):
    match = re.search(r"\((\d+)\s+months of data\)", str(title), flags=re.IGNORECASE)
    return int(match.group(1)) if match else None


def _choose_quarterly_series(claims_full, claims_drop, code_col, title):
    full = poisson_deck.build_ts(claims_full, code_col)
    dropped = poisson_deck.build_ts(claims_drop, code_col)
    target_n = _n_obs_hint(title)
    if target_n is not None:
        if int(dropped["y"].notna().sum()) == target_n:
            return dropped
        if int(full["y"].notna().sum()) == target_n:
            return full
    return dropped if not dropped.empty else full


def _rates_for_series(ts, enrollment, label):
    work = ts.copy()
    work["month"] = pd.to_datetime(work["month"]).dt.to_period("M").dt.to_timestamp()
    enr = enrollment.copy()
    enr["month"] = pd.to_datetime(enr["month"]).dt.to_period("M").dt.to_timestamp()
    work = work.merge(enr[["month", "enrollment"]], on="month", how="left", validate="many_to_one")
    if work["enrollment"].isna().any():
        missing = work.loc[work["enrollment"].isna(), "month"].dt.strftime("%Y-%m").tolist()
        raise ValueError(f"Missing enrollment for {label}: {missing}")
    return poisson_deck.utilization_rates_by_period(work, label, rate_per=RATE_PER)


# ---- build mapping target_title -> {quarter df, year df} -------------------

def compute_rate_tables():
    claims_full = poisson_deck.load_claims_pivot(CLAIMS_PATH)
    enrollment = poisson_deck.load_enrollment(ENROLLMENT_PATH)
    claims_drop = claims_full[
        ~claims_full.index.isin(
            pd.date_range(pd.Timestamp(DROP_MONTHS_FOR_ROBUST[0]),
                          pd.Timestamp(DROP_MONTHS_FOR_ROBUST[1]), freq="D")
        )
    ].copy()

    prs = Presentation(str(SOURCE_PPTX))
    slides = list(prs.slides)
    agg_lookup = poisson_deck.aggregate_title_lookup(claims_full)

    by_title = {}
    skipped = []
    for idx, slide in enumerate(slides):
        # drive off the quarterly slide itself so each table matches its own chart
        base_title = None
        for sh in slide.shapes:
            if sh.has_text_frame and _single_line(sh.text_frame.text).startswith(QUARTERLY_PREFIX):
                base_title = _single_line(sh.text_frame.text)[len(QUARTERLY_PREFIX):].strip()
                break
        if base_title is None:
            continue
        title = base_title
        try:
            code_col, target_title, _ = _claim_column_from_base(slide, idx + 1, base_title, claims_full, agg_lookup)
            if code_col not in claims_drop.columns and code_col in claims_full.columns:
                claims_drop[code_col] = claims_full.loc[claims_drop.index, code_col]
            ts = _choose_quarterly_series(claims_full, claims_drop, code_col, target_title)
            if ts.empty:
                raise ValueError("empty series")
            rates = _rates_for_series(ts, enrollment, target_title)
            q = rates[rates["period_type"].eq("quarter")].reset_index(drop=True)
            y = rates[rates["period_type"].eq("year")].reset_index(drop=True)
            by_title[poisson_deck.norm_text(target_title)] = {
                "target_title": target_title,
                "n_months": int(ts["y"].notna().sum()),
                "quarter": q,
                "year": y,
            }
        except Exception as exc:
            skipped.append((idx + 1, title, str(exc)))
    return by_title, skipped


def _fmt_rate(v):
    v = float(v)
    if v == 0:
        return "0"
    if v < 0.01:
        return f"{v:.4f}"
    if v < 1:
        return f"{v:.3f}"
    return f"{v:.2f}"


def _set_cell(cell, text, *, bold=False, size=8.5, fill=None, align="center"):
    from pptx.util import Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    cell.margin_left = cell.margin_right = Pt(2)
    cell.margin_top = cell.margin_bottom = Pt(1)
    tf = cell.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER if align == "center" else PP_ALIGN.LEFT
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = "Aptos"
    if fill is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(*fill)


def add_yearly_tables(source_pptx, output_pptx, by_title,
                      intervention_year=2021):
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor

    prs = Presentation(str(source_pptx))
    HEAD_FILL = (0x26, 0x38, 0x4D)       # dark slate (matches bar edge color)
    POST_FILL = (0xF2, 0xDE, 0xDE)       # light red tint for post-intervention years
    PRE_FILL = (0xEE, 0xF2, 0xF7)        # light blue tint for pre-intervention years

    added, unmatched = 0, []
    for slide in prs.slides:
        title_text = None
        for sh in slide.shapes:
            if sh.has_text_frame and _single_line(sh.text_frame.text).startswith(QUARTERLY_PREFIX):
                title_text = _single_line(sh.text_frame.text)[len(QUARTERLY_PREFIX):].strip()
                break
        if title_text is None:
            continue
        key = poisson_deck.norm_text(title_text)
        info = by_title.get(key)
        if info is None:
            unmatched.append(title_text)
            continue

        ydf = info["year"].copy()
        ydf = ydf[pd.to_numeric(ydf["exposure_months"], errors="coerce") > 0]
        years = ydf["period_label"].tolist()
        n = len(years)
        if n == 0:
            unmatched.append(title_text)
            continue

        # caption
        cap = slide.shapes.add_textbox(Inches(0.48), Inches(6.52),
                                       prs.slide_width - Inches(0.96), Inches(0.22))
        cp = cap.text_frame.paragraphs[0]
        cp.text = "Yearly rate per 100,000 member-months (exact 95% Poisson CIs):"
        cp.font.name = "Aptos"
        cp.font.bold = True
        cp.font.size = Pt(9)

        # table: 3 rows (Year / Rate / 95% CI) x (1 label + n year columns)
        rows, cols = 3, n + 1
        tbl_w = prs.slide_width - Inches(0.96)
        gtbl = slide.shapes.add_table(rows, cols, Inches(0.48), Inches(6.74),
                                      tbl_w, Inches(0.70)).table
        gtbl.first_row = False
        gtbl.horz_banding = False

        label_w = Inches(1.30)
        year_w = int((tbl_w - label_w) / n)
        gtbl.columns[0].width = label_w
        for c in range(1, cols):
            gtbl.columns[c].width = year_w
        for r in range(rows):
            gtbl.rows[r].height = Inches(0.70 / 3)

        _set_cell(gtbl.cell(0, 0), "Year", bold=True, size=8.5,
                  fill=HEAD_FILL, align="left")
        _set_cell(gtbl.cell(1, 0), "Rate", bold=True, size=8.5, align="left")
        _set_cell(gtbl.cell(2, 0), "95% CI", bold=True, size=8.5, align="left")
        # make header label text white
        gtbl.cell(0, 0).text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        for j, (_, row) in enumerate(ydf.iterrows(), start=1):
            yr = int(row["period_label"])
            post = yr >= intervention_year
            hfill = POST_FILL if post else PRE_FILL
            _set_cell(gtbl.cell(0, j), str(yr), bold=True, size=8.5, fill=hfill)
            _set_cell(gtbl.cell(1, j), _fmt_rate(row["rate_per_100000_member_months"]),
                      size=8.5)
            ci = f"{_fmt_rate(row['poisson_ci_lower_per_100000'])}–{_fmt_rate(row['poisson_ci_upper_per_100000'])}"
            _set_cell(gtbl.cell(2, j), ci, size=7.5)

        added += 1

    prs.save(str(output_pptx))
    return added, unmatched


if __name__ == "__main__":
    by_title, skipped = compute_rate_tables()
    print(f"computed rate tables for {len(by_title)} Normalized-Poisson slides")
    print(f"skipped: {len(skipped)}")
    for s in skipped:
        print("   SKIP", s)

    # cross-check against V111 CSV where month counts match
    csv = pd.read_csv("its_all_codes_reportRetsefV111_quarterly_rates.csv")
    csv_by = {poisson_deck.norm_text(t): g for t, g in csv.groupby("slide_title")}
    checked = mism = 0
    for key, info in by_title.items():
        if key not in csv_by:
            continue
        ours = info["quarter"][["period_label", "claims", "rate_per_100000_member_months"]].copy()
        theirs = csv_by[key][["period_label", "claims", "rate_per_100000_member_months"]].copy()
        m = ours.merge(theirs, on="period_label", suffixes=("_new", "_csv"))
        if len(m) != len(ours) or len(m) != len(theirs):
            continue  # different month coverage -> expected for the 75-vs-78 cases
        checked += 1
        if not (m["claims_new"].astype(int).equals(m["claims_csv"].astype(int))):
            mism += 1
            print("   CLAIMS MISMATCH:", info["target_title"])
    print(f"cross-checked {checked} slides against V111 csv, claim mismatches: {mism}")

    OUT = "its_all_codes_reportRetsefV12_with_quarterly_and_yearly_rates.pptx"
    added, unmatched = add_yearly_tables(SOURCE_PPTX, OUT, by_title)
    print(f"\nadded yearly tables to {added} quarterly slides -> {OUT}")
    print(f"unmatched quarterly slides: {len(unmatched)}")
    for u in unmatched:
        print("   UNMATCHED:", u)
