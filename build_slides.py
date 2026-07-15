"""
Build the Medicaid ITS slide deck from scratch.

For every target (an individual HCPCS code or a named custom aggregate) three
slides are produced, in order:

  1. Normalized Poisson ITS  - GLM Poisson with enrollment as the exposure
     offset (claims per 100,000 member-months). April 2022 enrollment is
     interpolated because the raw value is a reporting spike.
  2. Normalized Negative Binomial ITS - same design, NB2 likelihood.
  3. Quarterly Poisson rate bar chart (exact 95% CIs) plus a yearly
     summary table.

Edit the CONFIG block below to choose codes / aggregates and paths, then run:

    python build_slides.py

Input is claims_pivot.csv (built by Dataset_builder.ipynb) and
Enrollees_Merged.xlsx.  Nothing depends on a pre-existing template deck.
"""

import io
import warnings
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import statsmodels.api as sm
from matplotlib.ticker import FuncFormatter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from scipy.stats import chi2
from statsmodels.tools.sm_exceptions import ConvergenceWarning

# ============================ CONFIG ========================================

CLAIMS_PATH = "claims_pivot.csv"
ENROLLMENT_PATH = "Enrollees_Merged.xlsx"
OUTPUT_PPTX = "medicaid_its_deck.pptx"
UTILIZATION_RATES_CSV = "medicaid_its_utilization_rates.csv"

INTERVENTION_MONTH = "2021-03-01"          # ITS break point
RATE_PER = 100_000                          # claims per N member-months

# Enrollment months to replace with the average of two neighbours (reporting
# errors).  April 2022 is a known ~104M spike vs ~93M on either side.
ENROLLMENT_INTERPOLATION = {
    "2022-04-01": ("2022-03-01", "2022-05-01"),
}

# Individual HCPCS codes -> (code, short description).  One code = 3 slides.
INDIVIDUAL_CODES = [
    ("95943", "Quantitative autonomic function testing"),
    ("82785", "IgE (total)"),
    ("83520", "Tryptase"),
    ("82570", "Creatinine"),
    ("86160", "C3"),
    ("86162", "CH50"),
    ("86140", "CRP"),
    ("85652", "ESR"),
    ("86038", "ANA"),
    ("86235", "ENA panel"),
    ("86665", "EBV antibodies"),
    ("83690", "Lipase"),
    ("82150", "Amylase"),
    ("82784", "Total serum IgA"),
    ("83516", "Deamidated gliadin peptide antibodies"),
    ("83993", "Fecal calprotectin"),
    ("82274", "Fecal occult blood"),
    ("87045", "Stool culture"),
    ("87177", "Ova and parasites"),
    ("87338", "H. pylori stool antigen"),
    ("85390", "Coagulation interpretation/report"),
    ("85379", "D-dimer"),
    ("85384", "Fibrinogen"),
    ("75565", "Cardiac MRI"),
    ("75574", "Cardiac CTA"),
    ("93000", "ECG"),
    ("93303", "Complete transthoracic echocardiogram (TTE)"),
    ("93304", "TTE"),
    ("93307", "TTE"),
    ("93308", "TTE"),
    ("93306", "Echocardiogram"),
    ("85025", "CBC"),
]

# Named custom aggregates -> list of HCPCS codes summed into one series.
AGGREGATES = {
    "Custom Aggregate 95921 and 95923": ["95921", "95923"],
    "Custom Aggregate 93041 and 93042": ["93041", "93042"],
    "Custom Aggregate ECG 2": ["93042", "93000", "93005", "93010"],
    "Custom Aggregate EEG 1": ["95812", "95813", "95816"],
    "Custom Aggregate 93241-93248, 0295T-0298T, and 93224-93227": [
        "93241", "93242", "93243", "93244", "93245", "93246", "93247", "93248",
        "0295T", "0296T", "0297T", "0298T",
        "93224", "93225", "93226", "93227",
    ],
    "Aggregate Psychological Testing": [
        "96130", "96131", "96132", "96133", "96134", "96135",
        "96136", "96137", "96138", "96139", "96146",
    ],
    "Aggregate Skin Biopsy": [
        "11100", "11101", "11102", "11103", "11104", "11105", "11106", "11107",
    ],
    "Custom Aggregate EMG 3": [
        "95885", "95886", "95887",
        "95907", "95908", "95909", "95910", "95911", "95912", "95913",
    ],
    "Custom Aggregate 86146, 86147 and 86148": ["86146", "86147", "86148"],
    "Custom Aggregate 85613 and 83516": ["85613", "83516"],
    "Custom Aggregate Initial Hemostasis / Thrombotic Evaluation": [
        "85610", "85730", "85379", "85384",
    ],
    "Custom Aggregate Platelet Activation / Primary Hemostasis": ["85576", "85597"],
    "Custom Aggregate Thrombophilia Evaluation": [
        "85300", "85303", "85305", "85306", "81241", "81240",
    ],
    "Aggregate Clotting": [str(c) for c in range(85240, 85271)],
    "Custom Aggregate 71250, 71260, 71270 and 71271": ["71250", "71260", "71270", "71271"],
    "Custom Aggregate 75557 and 75561": ["75561", "75557"],
    "Custom Aggregate 93303, 93304, 93306, 93307 and 93308": [
        "93303", "93304", "93306", "93307", "93308",
    ],
    "Aggregate Fractures": [
        "23500", "23600", "25600", "27500", "27750", "27786", "28470", "21310",
        "28510", "25605", "24670", "27236", "24546", "25606", "27235",
    ],
}

# ============================ DATA LOADING ==================================

UTILIZATION_RATE_COLUMNS = [
    "code", "period_type", "period_label", "claims", "exposure_months",
    "rate_per_100000_member_months",
    "poisson_ci_lower_per_100000", "poisson_ci_upper_per_100000",
]


def load_claims_pivot(path):
    claims = pd.read_csv(path)
    date_col = claims.columns[0]
    claims[date_col] = pd.to_datetime(claims[date_col])
    return claims.set_index(date_col).sort_index()


def load_enrollment(path):
    enrollment = pd.read_excel(path)
    enrollment = enrollment.rename(columns={c: str(c).strip() for c in enrollment.columns})
    if {"Date", "Total Enrollees"}.issubset(enrollment.columns):
        enrollment = enrollment[["Date", "Total Enrollees"]].rename(
            columns={"Date": "month", "Total Enrollees": "enrollment"}
        )
    elif {"Month", "Enrollment"}.issubset(enrollment.columns):
        enrollment = enrollment[["Month", "Enrollment"]].rename(
            columns={"Month": "month", "Enrollment": "enrollment"}
        )
    else:
        raise ValueError("Enrollment file needs Date/Total Enrollees or Month/Enrollment columns.")

    enrollment["month"] = pd.to_datetime(enrollment["month"]).dt.to_period("M").dt.to_timestamp()
    enrollment["enrollment"] = pd.to_numeric(enrollment["enrollment"], errors="coerce")
    enrollment = enrollment.dropna(subset=["month", "enrollment"])
    enrollment = enrollment[enrollment["enrollment"] > 0]
    enrollment["enrollment"] = enrollment["enrollment"].astype(float)
    enrollment = enrollment.drop_duplicates(subset=["month"], keep="last").sort_values("month")

    values = enrollment.set_index("month")["enrollment"]
    for target, (prev, nxt) in ENROLLMENT_INTERPOLATION.items():
        target, prev, nxt = (pd.Timestamp(x) for x in (target, prev, nxt))
        if {target, prev, nxt}.issubset(values.index):
            interp = (values.loc[prev] + values.loc[nxt]) / 2
            enrollment.loc[enrollment["month"] == target, "enrollment"] = interp
    return enrollment


def build_series(claims, codes):
    """Sum the given code columns into one monthly series, trimmed to the
    first month with a non-zero count."""
    lookup = {str(c).strip().upper(): c for c in claims.columns}
    present = [lookup[str(c).strip().upper()] for c in codes if str(c).strip().upper() in lookup]
    if not present:
        raise ValueError(f"none of the codes are present in the claims pivot: {codes}")
    y = claims[present].sum(axis=1, min_count=1)
    ts = pd.DataFrame(
        {"month": pd.to_datetime(claims.index), "y": pd.to_numeric(y, errors="coerce")}
    ).reset_index(drop=True)
    nonzero = ts.index[ts["y"].fillna(0) > 0]
    if len(nonzero) == 0:
        return ts.iloc[0:0].copy()
    return ts.loc[nonzero[0]:].reset_index(drop=True)

# ============================ MODELS ========================================

def add_its_terms(ts, enrollment, intervention_month):
    ts = ts.copy()
    ts["month"] = pd.to_datetime(ts["month"]).dt.to_period("M").dt.to_timestamp()
    ts["y"] = pd.to_numeric(ts["y"], errors="coerce").fillna(0)

    enr = enrollment.copy()
    enr["month"] = pd.to_datetime(enr["month"]).dt.to_period("M").dt.to_timestamp()
    ts = ts.merge(enr[["month", "enrollment"]], on="month", how="left", validate="many_to_one")
    if ts["enrollment"].isna().any():
        missing = ts.loc[ts["enrollment"].isna(), "month"].dt.strftime("%Y-%m").tolist()
        raise ValueError(f"missing enrollment for months: {missing}")

    idx = ts["month"].dt.year * 12 + ts["month"].dt.month
    ts["T"] = idx - idx.iloc[0]
    intervention = pd.Timestamp(intervention_month).to_period("M").to_timestamp()
    intervention_idx = intervention.year * 12 + intervention.month
    ts["X_t"] = (ts["month"] >= intervention).astype(int)
    ts["post_time"] = 0
    post = ts["month"] >= intervention
    ts.loc[post, "post_time"] = idx[post] - intervention_idx + 1

    x = sm.add_constant(ts[["T", "X_t", "post_time"]])
    return ts, x


def fit_poisson(ts, enrollment, intervention_month, rate_per=RATE_PER):
    fitted, x = add_its_terms(ts, enrollment, intervention_month)
    model = sm.GLM(fitted["y"], x, family=sm.families.Poisson(), exposure=fitted["enrollment"]).fit()
    fitted["y_hat"] = model.predict(x, exposure=fitted["enrollment"])
    fitted["rate"] = fitted["y"] / fitted["enrollment"] * rate_per
    fitted["rate_hat"] = fitted["y_hat"] / fitted["enrollment"] * rate_per
    return fitted, model


def fit_negative_binomial(fitted_ts, intervention_month, rate_per=RATE_PER):
    ts, x = add_its_terms(fitted_ts[["month", "y"]], fitted_ts[["month", "enrollment"]], intervention_month)
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", category=ConvergenceWarning)
        warnings.filterwarnings("ignore", category=RuntimeWarning)
        model = sm.NegativeBinomial(
            ts["y"], x, exposure=ts["enrollment"], loglike_method="nb2"
        ).fit(disp=False, maxiter=200)
    ts["y_hat"] = model.predict(x, exposure=ts["enrollment"])
    ts["rate"] = ts["y"] / ts["enrollment"] * rate_per
    ts["rate_hat"] = ts["y_hat"] / ts["enrollment"] * rate_per
    return ts, model

# ============================ RATES / CIs ===================================

def exact_poisson_count_ci(count, alpha=0.05):
    count = int(count)
    if count < 0:
        raise ValueError("Poisson CI requires non-negative claims.")
    lower = 0.0 if count == 0 else 0.5 * chi2.ppf(alpha / 2, 2 * count)
    upper = 0.5 * chi2.ppf(1 - alpha / 2, 2 * (count + 1))
    return float(lower), float(upper)


def _period_rate_row(code, period_type, label, claims, exposure_months, rate_per=RATE_PER):
    claims = int(round(float(claims)))
    exposure_months = float(exposure_months)
    if exposure_months <= 0:
        raise ValueError(f"exposure must be positive for {code} {label}.")
    lo, hi = exact_poisson_count_ci(claims)
    return {
        "code": str(code),
        "period_type": period_type,
        "period_label": label,
        "claims": claims,
        "exposure_months": exposure_months,
        "rate_per_100000_member_months": claims / exposure_months * rate_per,
        "poisson_ci_lower_per_100000": lo / exposure_months * rate_per,
        "poisson_ci_upper_per_100000": hi / exposure_months * rate_per,
    }


def utilization_rates_by_period(ts, code, rate_per=RATE_PER):
    work = ts.copy()
    work["month"] = pd.to_datetime(work["month"]).dt.to_period("M").dt.to_timestamp()
    work["y"] = pd.to_numeric(work["y"], errors="coerce").fillna(0)
    work["enrollment"] = pd.to_numeric(work["enrollment"], errors="coerce")
    work = work.dropna(subset=["month", "enrollment"])
    work = work[work["enrollment"] > 0]

    specs = [
        ("month", work["month"].dt.to_period("M"), lambda p: p.strftime("%Y-%m")),
        ("quarter", work["month"].dt.to_period("Q"), lambda p: f"{p.year}-Q{p.quarter}"),
        ("year", work["month"].dt.to_period("Y"), lambda p: str(p.year)),
    ]
    rows = []
    for period_type, periods, label_of in specs:
        grouped = work.assign(_p=periods).groupby("_p", sort=True).agg(
            claims=("y", "sum"), exposure_months=("enrollment", "sum")
        )
        for period, row in grouped.iterrows():
            rows.append(_period_rate_row(code, period_type, label_of(period),
                                         row["claims"], row["exposure_months"], rate_per))
    return pd.DataFrame(rows, columns=UTILIZATION_RATE_COLUMNS)

# ============================ PLOTS =========================================

def plot_its(ts, title, intervention_month, fitted_label, rate_per=RATE_PER):
    fig, ax = plt.subplots(figsize=(10, 8))
    ax.plot(ts["month"], ts["rate"], marker="o", linewidth=1.5,
            label=f"Observed claims per {rate_per:,} member-months")
    ax.plot(ts["month"], ts["rate_hat"], linestyle="--", linewidth=2, label=fitted_label)
    ax.axvline(pd.Timestamp(intervention_month), color="red", linestyle=":", linewidth=2,
               label="Intervention")
    ax.set_xlabel("Month")
    ax.set_ylabel(f"Claims per {rate_per:,} Medicaid member-months")
    y_max = float(np.nanmax(np.r_[ts["rate"].to_numpy(float), ts["rate_hat"].to_numpy(float), [0.0]]))
    ax.set_ylim(0, y_max * 1.05 if y_max > 0 else 1.0)
    ax.ticklabel_format(axis="y", style="plain", useOffset=False)
    ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: f"{v:,.2f}".rstrip("0").rstrip(".")))
    ax.set_title(title)
    ax.legend()
    ax.grid(alpha=0.25)
    plt.xticks(rotation=45)
    plt.tight_layout()
    return fig


def plot_quarterly_bars(rates, title, rate_per=RATE_PER, intervention_quarter=None):
    df = rates[rates["period_type"].eq("quarter")].copy()
    for col in ("rate_per_100000_member_months",
                "poisson_ci_lower_per_100000", "poisson_ci_upper_per_100000"):
        df[col] = pd.to_numeric(df[col], errors="coerce")
    df = df.dropna(subset=["rate_per_100000_member_months"]).reset_index(drop=True)

    x = np.arange(len(df))
    y = df["rate_per_100000_member_months"].to_numpy(float)
    lower = np.maximum(y - df["poisson_ci_lower_per_100000"].to_numpy(float), 0)
    upper = np.maximum(df["poisson_ci_upper_per_100000"].to_numpy(float) - y, 0)

    fig_width = max(11, min(18, 0.36 * max(len(df), 1)))
    fig, ax = plt.subplots(figsize=(fig_width, 6.1))
    ax.bar(x, y, color="#4C78A8", edgecolor="#26384D", linewidth=0.4,
           yerr=np.vstack([lower, upper]),
           error_kw={"ecolor": "#222222", "elinewidth": 0.9, "capsize": 2.5, "capthick": 0.9})
    ax.set_title(title)
    ax.set_xlabel("Quarter")
    ax.set_ylabel(f"Claims per {rate_per:,} member-months")
    ax.set_xticks(x)
    ax.set_xticklabels(df["period_label"], rotation=65, ha="right")
    ax.ticklabel_format(axis="y", style="plain", useOffset=False)
    ax.grid(axis="y", alpha=0.25)
    ax.set_axisbelow(True)
    if intervention_quarter and intervention_quarter in set(df["period_label"]):
        pos = df.index[df["period_label"].eq(intervention_quarter)][0]
        ax.axvline(pos - 0.5, color="#C44E52", linestyle=":", linewidth=1.6)
    y_max = float(np.nanmax(np.r_[y + upper, [0.0]]))
    ax.set_ylim(0, y_max * 1.10 if y_max > 0 else 1.0)
    fig.tight_layout()
    return fig


def _fig_to_stream(fig, dpi=200):
    stream = io.BytesIO()
    fig.savefig(stream, format="png", dpi=dpi, bbox_inches="tight")
    plt.close(fig)
    stream.seek(0)
    return stream

# ============================ SLIDE HELPERS =================================

def clean_summary_text(summary_text):
    text = str(summary_text).replace("\r\n", "\n")
    marker = "\nNotes:\n"
    if marker in text:
        text = text.split(marker, 1)[0].rstrip()
    return text


def pick_font_size(text, box_width_pts, box_height_pts):
    lines = text.splitlines() or [""]
    max_len = max(len(line) for line in lines)
    for size in (12, 11, 10, 9, 8, 7, 6):
        if max_len <= box_width_pts / (0.60 * size) and len(lines) <= box_height_pts / (1.20 * size):
            return size
    return 6


def _add_title(slide, text, prs):
    box = slide.shapes.add_textbox(Inches(0.35), Inches(0.2), prs.slide_width - Inches(0.7), Inches(0.62))
    p = box.text_frame.paragraphs[0]
    p.text = text
    p.font.name = "Aptos"
    p.font.bold = True
    p.font.size = Pt(22)


def _add_codes_note(slide, codes, prs, top):
    box = slide.shapes.add_textbox(Inches(0.35), top, prs.slide_width - Inches(0.7), Inches(0.5))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Codes in aggregation: " + ", ".join(str(c) for c in codes)
    p.font.name = "Aptos"
    p.font.size = Pt(8)


def add_model_slide(prs, slide_title, ts, model, intervention_month, fitted_label,
                    rate_per=RATE_PER, agg_codes=None, error=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_title(slide, slide_title, prs)

    if model is not None:
        fig = plot_its(ts, slide_title, intervention_month, fitted_label, rate_per=rate_per)
        slide.shapes.add_picture(_fig_to_stream(fig), Inches(0.35), Inches(0.95),
                                 Inches(7.0), Inches(5.9))
        summary_text = clean_summary_text(model.summary().as_text())
    else:
        summary_text = "Model could not be fit:\n" + str(error)

    box_left, box_top = Inches(7.55), Inches(1.0)
    box_width, box_height = Inches(5.4), Inches(5.6)
    summary_box = slide.shapes.add_textbox(box_left, box_top, box_width, box_height)
    sf = summary_box.text_frame
    sf.word_wrap = False
    sf.margin_left = sf.margin_right = sf.margin_top = sf.margin_bottom = 0
    sf.text = summary_text
    size = pick_font_size(summary_text, box_width / 12700, box_height / 12700)
    for para in sf.paragraphs:
        para.font.name = "Consolas"
        para.font.size = Pt(size)

    if agg_codes:
        _add_codes_note(slide, agg_codes, prs, prs.slide_height - Inches(0.6))
    return slide


def _set_cell(cell, text, *, bold=False, size=8.5, fill=None, align="center", white=False):
    cell.margin_left = cell.margin_right = Pt(2)
    cell.margin_top = cell.margin_bottom = Pt(1)
    tf = cell.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER if align == "center" else PP_ALIGN.LEFT
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = "Aptos"
    if white:
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    if fill is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(*fill)


def _fmt_rate(v):
    v = float(v)
    if v == 0:
        return "0"
    if v < 0.01:
        return f"{v:.4f}"
    if v < 1:
        return f"{v:.3f}"
    return f"{v:.2f}"


def _add_yearly_table(slide, year_df, prs, intervention_year):
    HEAD_FILL, POST_FILL, PRE_FILL = (0x26, 0x38, 0x4D), (0xF2, 0xDE, 0xDE), (0xEE, 0xF2, 0xF7)
    ydf = year_df[pd.to_numeric(year_df["exposure_months"], errors="coerce") > 0]
    n = len(ydf)
    if n == 0:
        return

    cap = slide.shapes.add_textbox(Inches(0.48), Inches(6.52), prs.slide_width - Inches(0.96), Inches(0.22))
    cp = cap.text_frame.paragraphs[0]
    cp.text = "Yearly rate per 100,000 member-months (exact 95% Poisson CIs):"
    cp.font.name = "Aptos"
    cp.font.bold = True
    cp.font.size = Pt(9)

    tbl_w = prs.slide_width - Inches(0.96)
    table = slide.shapes.add_table(3, n + 1, Inches(0.48), Inches(6.74), tbl_w, Inches(0.70)).table
    table.first_row = False
    table.horz_banding = False
    label_w = Inches(1.30)
    year_w = int((tbl_w - label_w) / n)
    table.columns[0].width = label_w
    for c in range(1, n + 1):
        table.columns[c].width = year_w
    for r in range(3):
        table.rows[r].height = Inches(0.70 / 3)

    _set_cell(table.cell(0, 0), "Year", bold=True, fill=HEAD_FILL, align="left", white=True)
    _set_cell(table.cell(1, 0), "Rate", bold=True, align="left")
    _set_cell(table.cell(2, 0), "95% CI", bold=True, align="left")
    for j, (_, row) in enumerate(ydf.iterrows(), start=1):
        yr = int(row["period_label"])
        hfill = POST_FILL if yr >= intervention_year else PRE_FILL
        _set_cell(table.cell(0, j), str(yr), bold=True, fill=hfill)
        _set_cell(table.cell(1, j), _fmt_rate(row["rate_per_100000_member_months"]))
        ci = f"{_fmt_rate(row['poisson_ci_lower_per_100000'])}–{_fmt_rate(row['poisson_ci_upper_per_100000'])}"
        _set_cell(table.cell(2, j), ci, size=7.5)


def add_quarterly_slide(prs, target_title, rates, rate_per=RATE_PER,
                        intervention_quarter=None, intervention_year=2021, agg_codes=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title = f"Quarterly rate per {rate_per:,} member-months: {target_title}"
    _add_title(slide, title, prs)

    fig = plot_quarterly_bars(rates, title, rate_per=rate_per, intervention_quarter=intervention_quarter)
    slide.shapes.add_picture(_fig_to_stream(fig, dpi=220), Inches(0.30), Inches(0.90),
                             prs.slide_width - Inches(0.60), Inches(5.35))

    note_parts = ["Error bars show exact 95% Poisson confidence intervals."]
    if agg_codes:
        note_parts.append("Codes in aggregation: " + ", ".join(str(c) for c in agg_codes))
    note = slide.shapes.add_textbox(Inches(0.48), Inches(6.28), prs.slide_width - Inches(0.96), Inches(0.22))
    np_ = note.text_frame.paragraphs[0]
    np_.text = " ".join(note_parts)
    np_.font.name = "Aptos"
    np_.font.size = Pt(8)

    _add_yearly_table(slide, rates[rates["period_type"].eq("year")].reset_index(drop=True),
                      prs, intervention_year)
    return slide

# ============================ DRIVER ========================================

def _targets():
    """Yield (title, code, code_list, is_aggregate) for every configured target."""
    for code, desc in INDIVIDUAL_CODES:
        title = f"HCPCS {code}" + (f" – {desc}" if desc else "")
        yield title, code, [code], False
    for name, codes in AGGREGATES.items():
        yield name, name, list(codes), True


def build_deck():
    claims = load_claims_pivot(CLAIMS_PATH)
    enrollment = load_enrollment(ENROLLMENT_PATH)

    intervention = pd.Timestamp(INTERVENTION_MONTH)
    intervention_quarter = f"{intervention.year}-Q{(intervention.month - 1) // 3 + 1}"

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    rate_frames = []
    skipped = []
    made = 0
    for title, code, codes, is_agg in _targets():
        agg_codes = codes if is_agg else None
        try:
            ts = build_series(claims, codes)
            if ts.empty:
                raise ValueError("no non-zero observations")

            fitted, poisson_model = fit_poisson(ts, enrollment, intervention, rate_per=RATE_PER)
            rates = utilization_rates_by_period(fitted, code, rate_per=RATE_PER)
            rates.insert(0, "target_title", title)
            rate_frames.append(rates)

            nb_ts, nb_model, nb_error = fitted, None, None
            try:
                nb_ts, nb_model = fit_negative_binomial(fitted, intervention, rate_per=RATE_PER)
            except Exception as exc:  # NB can fail to converge for sparse series
                nb_error = exc

            add_model_slide(
                prs, f"Normalized Poisson ITS (per {RATE_PER:,} member-months): {title}",
                fitted, poisson_model, intervention,
                f"Fitted Poisson ITS rate per {RATE_PER:,} member-months",
                rate_per=RATE_PER, agg_codes=agg_codes,
            )
            add_model_slide(
                prs, f"Normalized Negative Binomial ITS (per {RATE_PER:,} member-months): {title}",
                nb_ts, nb_model, intervention,
                f"Fitted Negative Binomial ITS rate per {RATE_PER:,} member-months",
                rate_per=RATE_PER, agg_codes=agg_codes, error=nb_error,
            )
            add_quarterly_slide(
                prs, title, rates, rate_per=RATE_PER,
                intervention_quarter=intervention_quarter,
                intervention_year=intervention.year, agg_codes=agg_codes,
            )
            made += 1
        except Exception as exc:
            skipped.append((title, str(exc)))

    Path(OUTPUT_PPTX).parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT_PPTX)

    rates_out = (pd.concat(rate_frames, ignore_index=True)
                 if rate_frames else pd.DataFrame(columns=["target_title"] + UTILIZATION_RATE_COLUMNS))
    rates_out.to_csv(UTILIZATION_RATES_CSV, index=False)

    return {
        "output_pptx": OUTPUT_PPTX,
        "utilization_rates_csv": UTILIZATION_RATES_CSV,
        "targets_rendered": made,
        "slides_created": made * 3,
        "skipped": skipped,
    }


if __name__ == "__main__":
    report = build_deck()
    print(f"Deck:  {report['output_pptx']}")
    print(f"Rates: {report['utilization_rates_csv']}")
    print(f"Targets rendered: {report['targets_rendered']}  (slides: {report['slides_created']})")
    if report["skipped"]:
        print(f"Skipped {len(report['skipped'])}:")
        for title, reason in report["skipped"]:
            print(f"   - {title}: {reason}")
